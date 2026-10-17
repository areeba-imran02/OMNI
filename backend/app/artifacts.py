import io, re
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document

def slug(s):
    out = re.sub(r"[^a-zA-Z0-9]+", "-", s).strip("-").lower()
    return (out[:50] or "omnicraft-result")

def make_markdown_artifact(title, content):
    return {"name": slug(title) + ".md", "type": "markdown", "content": f"# {title}\n\n{content}"}

def escape_html(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def make_pdf(title, content):
    buff = io.BytesIO()
    doc = SimpleDocTemplate(buff, pagesize=letter, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()
    story = [Paragraph(escape_html(title), styles["Title"]), Spacer(1, 16)]
    for block in content.split("\n\n"):
        if block.strip().startswith("```"):
            story.append(Preformatted(escape_html(block.replace("```", "")), styles["Code"]))
        else:
            story.append(Paragraph(escape_html(block).replace("\n", "<br/>"), styles["BodyText"]))
            story.append(Spacer(1, 8))
    doc.build(story)
    return buff.getvalue()

def make_docx(title, content):
    doc = Document()
    doc.add_heading(title, 0)
    for line in content.splitlines():
        if line.startswith("# "): doc.add_heading(line[2:], 1)
        elif line.startswith("## "): doc.add_heading(line[3:], 2)
        elif line.startswith("- "): doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip(): doc.add_paragraph(line)
    buff = io.BytesIO()
    doc.save(buff)
    return buff.getvalue()
